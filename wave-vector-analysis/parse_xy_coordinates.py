#!/usr/bin/env python3
"""
Parse the manual `XY coordinates.xlsx` ground-truth into a tidy CSV.

Each sheet (one per video) holds two blocks: a "Background video" block and a
"Calcium video" block. We keep the CALCIUM block because its pixel coordinates
are in the same space as the C-channel TIFFs used by the wave catalog.

Output columns (one row per annotated point):
  prefix          - video id (WG30, PS63, WGD81, WS61, ...)
  video_stem      - calcium TIFF stem when available (else "")
  side            - Left / Right / single   (which embryo)
  landmark_raw    - verbatim label
  landmark_class  - head | tail | poke | pressure | cement_gland | eye |
                    tail_response | local | other
  frame           - annotated frame (int or "")
  x, y            - pixel coordinates

Run:
  python parse_xy_coordinates.py \
    --xlsx "/path/to/XY coordinates.xlsx" \
    --out analysis_results/xy_ground_truth.csv
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
from provenance import csv_comment_header, record_run  # noqa: E402
DEFAULT_XLSX = Path(
    "/Users/jdietz/Library/CloudStorage/Box-Box/Calcium videos/XY coordinates.xlsx"
)

LANDMARK_CLASS = [
    (r"^head$", "head"),
    (r"^tail$", "tail"),
    (r"poke", "poke"),
    (r"pressure", "pressure"),
    (r"cement gland|^cg ", "cement_gland"),
    (r"\beye\b", "eye"),
    (r"tail (response|signal)", "tail_response"),
    (r"local signal", "local"),
]


def classify_landmark(label: str) -> str:
    s = label.strip().lower()
    for pat, cls in LANDMARK_CLASS:
        if re.search(pat, s):
            return cls
    return "other"


def _is_tif(s) -> bool:
    return isinstance(s, str) and s.strip().lower().endswith((".tif", ".tiff"))


def parse_sheet(ws, sheet_name: str) -> list[dict]:
    prefix = re.sub(r"^Layer\s*\d+\s*-\s*", "", sheet_name).strip()
    rows = list(ws.iter_rows(values_only=True))

    section = None          # 'background' | 'calcium'
    calcium_stem = ""
    header = None
    col = {}
    out = []

    def set_header(r):
        nonlocal header, col
        header = [str(c).strip().lower() if c else "" for c in r]
        col = {name: header.index(name) for name in
               ("frame", "orientation", "location", "x (pixel)", "y (pixel)")
               if name in header}

    for r in rows:
        cells = list(r)
        first_str = next((str(c).strip() for c in cells
                          if isinstance(c, str) and str(c).strip()), "")

        if first_str in ("Background video", "Calcium video"):
            section = "background" if first_str == "Background video" else "calcium"
            header = None
            continue
        if first_str == "Results":
            section = None
            header = None
            continue

        # capture calcium TIFF name
        if section == "calcium" and not calcium_stem:
            tif = next((c for c in cells if _is_tif(c)), None)
            if tif:
                calcium_stem = Path(str(tif).strip()).stem

        # header row?
        if any(isinstance(c, str) and c.strip().lower() == "frame" for c in cells):
            set_header(cells)
            continue

        # data row in the calcium block only
        if section == "calcium" and header and "location" in col:
            li = col["location"]
            if li >= len(cells) or not cells[li]:
                continue
            label = str(cells[li]).strip()
            if not label:
                continue
            oi = col.get("orientation")
            xi = col.get("x (pixel)")
            yi = col.get("y (pixel)")
            fi = col.get("frame")
            side = ""
            if oi is not None and oi < len(cells) and cells[oi]:
                side = str(cells[oi]).strip().capitalize()
            side = side or "single"
            frame = cells[fi] if (fi is not None and fi < len(cells)) else ""
            x = cells[xi] if (xi is not None and xi < len(cells)) else ""
            y = cells[yi] if (yi is not None and yi < len(cells)) else ""
            if x == "" or y == "":
                continue
            out.append({
                "prefix": prefix,
                "video_stem": calcium_stem,
                "side": side,
                "landmark_raw": label,
                "landmark_class": classify_landmark(label),
                "frame": int(frame) if isinstance(frame, (int, float)) else "",
                "x": float(x), "y": float(y),
            })
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--out",
                    default=str(HERE / "analysis_results" / "xy_ground_truth.csv"))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    all_rows = []
    for sn in wb.sheetnames:
        all_rows.extend(parse_sheet(wb[sn], sn))

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    fields = ["prefix", "video_stem", "side", "landmark_raw",
              "landmark_class", "frame", "x", "y"]
    rec = record_run(
        "parse_xy_coordinates.py", out,
        inputs={"xlsx": args.xlsx},
        extra={
            "n_points": len(all_rows),
            "n_videos": len(set(r["prefix"] for r in all_rows)),
        },
    )
    with open(out, "w", newline="") as f:
        f.write(csv_comment_header(rec))
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(all_rows)

    # console summary
    from collections import Counter
    by_class = Counter(r["landmark_class"] for r in all_rows)
    print(f"Wrote {out}  ({len(all_rows)} points, "
          f"{len(set(r['prefix'] for r in all_rows))} videos)")
    for cls, n in by_class.most_common():
        vids = len(set(r["prefix"] for r in all_rows
                       if r["landmark_class"] == cls))
        print(f"  {cls:14} {n:4} points  in {vids} videos")


if __name__ == "__main__":
    main()
